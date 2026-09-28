import csv
import hashlib
import io
import re
import uuid

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.schemas import InvoiceOut, InvoiceUpdate
from app.core.auth import current_user_id
from app.core.config import get_settings
from app.db.session import get_db
from app.models import Invoice
from app.services.pipeline import apply_extraction
from app.services.validation import ExtractedInvoice

router = APIRouter(prefix="/invoices", tags=["invoices"])

ALLOWED_TYPES = {
    "application/pdf": ".pdf",
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/webp": ".webp",
}
MAGIC = {
    "application/pdf": (b"%PDF",),
    "image/png": (b"\x89PNG",),
    "image/jpeg": (b"\xff\xd8\xff",),
    "image/webp": (b"RIFF",),
}


def _get_owned(db: Session, invoice_id: uuid.UUID, user_id: uuid.UUID) -> Invoice:
    inv = db.get(Invoice, invoice_id)
    if inv is None or inv.user_id != user_id:  # same 404 for "not yours" and "not there"
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Invoice not found")
    return inv


def _safe_name(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", name)[:120] or "file"


@router.post("", response_model=list[InvoiceOut], status_code=status.HTTP_201_CREATED)
def upload(
    request: Request,
    files: list[UploadFile] = File(...),
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    # A plain `def`, not `async def`: this does blocking I/O (sync DB calls, and — for the
    # Supabase storage backend — a real HTTP POST per file). FastAPI/Starlette runs sync path
    # functions in a worker thread pool automatically; an `async def` version of this same
    # code would block the single event loop for its entire duration, freezing every other
    # request on the server (of any user) for as long as this upload takes. Every other
    # endpoint in this file is already a plain `def` for the same reason — this was the one
    # exception, found by that exact symptom: uploading several files made /health itself
    # stop responding.
    settings = get_settings()
    limit = settings.max_upload_mb * 1024 * 1024
    if not 1 <= len(files) <= 20:
        raise HTTPException(400, "Upload between 1 and 20 files at a time")

    prepared = []
    seen_hashes: dict[str, str] = {}  # hash -> filename, catches duplicates within this batch
    for f in files:
        data = f.file.read(limit + 1)  # sync read; see the note on `def upload` above
        if len(data) > limit:
            raise HTTPException(413, f"{f.filename}: larger than {settings.max_upload_mb} MB")
        ctype = f.content_type or ""
        if ctype not in ALLOWED_TYPES or not data.startswith(MAGIC[ctype]):
            raise HTTPException(415, f"{f.filename}: only PDF, PNG, JPEG or WebP files")

        file_hash = hashlib.sha256(data).hexdigest()
        if file_hash in seen_hashes:
            raise HTTPException(
                409, f"{f.filename}: identical to {seen_hashes[file_hash]} in this same upload"
            )
        seen_hashes[file_hash] = f.filename or "file"
        prepared.append((f.filename or "file", ctype, data, file_hash))

    # Reject the whole batch if any file's exact bytes were already uploaded by this user —
    # cheaper and clearer than letting it burn an extraction call only to be flagged after.
    existing = db.scalars(
        select(Invoice).where(Invoice.user_id == user_id, Invoice.file_hash.in_(seen_hashes))
    ).all()
    if existing:
        dupe = existing[0]
        raise HTTPException(
            409,
            f"This file was already uploaded as \"{dupe.filename}\" ({dupe.status}). "
            "Delete that one first, or use Retry on it, instead of uploading it again.",
        )

    storage, queue = request.app.state.storage, request.app.state.queue
    created = []
    for name, ctype, data, file_hash in prepared:
        inv_id = uuid.uuid4()
        path = f"{user_id}/{inv_id}{ALLOWED_TYPES[ctype]}"
        storage.save(path, data, ctype)
        inv = Invoice(id=inv_id, user_id=user_id, filename=_safe_name(name),
                      storage_path=path, content_type=ctype, status="queued", issues=[],
                      file_hash=file_hash)
        db.add(inv)
        db.flush()
        queue.enqueue(inv.id, db)
        created.append(inv)
    db.commit()
    return created


@router.get("", response_model=list[InvoiceOut])
def list_invoices(
    status_filter: str | None = None,
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    q = select(Invoice).where(Invoice.user_id == user_id).order_by(Invoice.created_at.desc())
    if status_filter:
        q = q.where(Invoice.status == status_filter)
    return db.scalars(q.limit(200)).all()


EXPORT_HEADERS = ["vendor", "invoice_number", "date", "currency", "subtotal", "tax", "total"]


def _approved_invoices(db: Session, user_id: uuid.UUID) -> list[Invoice]:
    return db.scalars(
        select(Invoice)
        .where(Invoice.user_id == user_id, Invoice.status == "approved")
        .order_by(Invoice.invoice_date, Invoice.created_at)
    ).all()


def _guard_formula(value) -> str:
    text = "" if value is None else str(value)
    # Spreadsheet formula injection: a vendor named "=HYPERLINK(...)" must stay text. Numbers
    # are never passed through this — they may legitimately be negative.
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


@router.get("/export.csv")
def export_csv(
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(EXPORT_HEADERS)
    for inv in _approved_invoices(db, user_id):
        w.writerow([
            _guard_formula(inv.vendor), _guard_formula(inv.invoice_number), inv.invoice_date,
            _guard_formula(inv.currency), inv.subtotal, inv.tax, inv.total,
        ])
    return Response(
        out.getvalue(), media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="invoices.csv"'},
    )


@router.get("/export.xlsx")
def export_xlsx(
    db: Session = Depends(get_db),
    user_id: uuid.UUID = Depends(current_user_id),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Approved invoices"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    text_cols = {1, 2, 4}  # vendor, invoice_number, currency — force-text, same guard as CSV
    for inv in _approved_invoices(db, user_id):
        row = [
            _guard_formula(inv.vendor), _guard_formula(inv.invoice_number),
            inv.invoice_date.isoformat() if inv.invoice_date else None,
            _guard_formula(inv.currency),
            float(inv.subtotal) if inv.subtotal is not None else None,
            float(inv.tax) if inv.tax is not None else None,
            float(inv.total) if inv.total is not None else None,
        ]
        ws.append(row)
        for col in text_cols:
            ws.cell(row=ws.max_row, column=col).number_format = "@"

    for col, width in zip("ABCDEFG", (24, 16, 12, 10, 12, 12, 12), strict=True):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="invoices.xlsx"'},
    )


@router.get("/{invoice_id}", response_model=InvoiceOut)
def get_invoice(invoice_id: uuid.UUID, db: Session = Depends(get_db),
                user_id: uuid.UUID = Depends(current_user_id)):
    return _get_owned(db, invoice_id, user_id)


@router.get("/{invoice_id}/file")
def get_file(invoice_id: uuid.UUID, request: Request, db: Session = Depends(get_db),
             user_id: uuid.UUID = Depends(current_user_id)):
    inv = _get_owned(db, invoice_id, user_id)
    return Response(request.app.state.storage.read(inv.storage_path),
                    media_type=inv.content_type,
                    headers={"Content-Disposition": "inline"})


@router.put("/{invoice_id}", response_model=InvoiceOut)
def review(invoice_id: uuid.UUID, body: InvoiceUpdate, db: Session = Depends(get_db),
           user_id: uuid.UUID = Depends(current_user_id)):
    inv = _get_owned(db, invoice_id, user_id)
    if inv.status in ("queued", "processing"):
        raise HTTPException(409, "Invoice is still being processed")

    data = ExtractedInvoice.model_validate(body.model_dump())
    apply_extraction(db, inv, data)  # re-runs validation + duplicate check on the edits
    if body.approve:
        if inv.issues and not body.acknowledge_issues:
            raise HTTPException(422, {"message": "Unresolved issues", "issues": inv.issues})
        inv.status = "approved"
    else:
        inv.status = "needs_review"
    db.commit()
    return inv


@router.post("/{invoice_id}/retry", response_model=InvoiceOut)
def retry(invoice_id: uuid.UUID, request: Request, db: Session = Depends(get_db),
          user_id: uuid.UUID = Depends(current_user_id)):
    inv = _get_owned(db, invoice_id, user_id)
    if inv.status != "failed":
        raise HTTPException(409, "Only failed invoices can be retried")
    inv.status, inv.error = "queued", None
    request.app.state.queue.enqueue(inv.id, db)
    db.commit()
    return inv


@router.delete("/{invoice_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete(invoice_id: uuid.UUID, request: Request, db: Session = Depends(get_db),
           user_id: uuid.UUID = Depends(current_user_id)):
    inv = _get_owned(db, invoice_id, user_id)
    path = inv.storage_path
    db.delete(inv)
    db.commit()
    request.app.state.storage.delete(path)
